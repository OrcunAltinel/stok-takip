"""Excel (.xlsx) aktarım yardımcısı — openpyxl."""

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def excel_aktar(
    baslik: str,
    sutunlar: list[str],
    satirlar: list[list[Any]],
    kayit_yolu: Path,
) -> None:
    """Genel amaçlı Excel aktarım fonksiyonu."""
    wb = Workbook()
    ws = wb.active
    ws.title = baslik[:31]

    # Başlık satırı
    baslik_font = Font(bold=True, color="FFFFFF")
    baslik_doldur = PatternFill(fill_type="solid", fgColor="1E3A5F")
    for col_idx, sutun in enumerate(sutunlar, 1):
        hucre = ws.cell(row=1, column=col_idx, value=sutun)
        hucre.font = baslik_font
        hucre.fill = baslik_doldur
        hucre.alignment = Alignment(horizontal="center")

    # Veri satırları
    zebra_doldur = PatternFill(fill_type="solid", fgColor="F0F4F8")
    for row_idx, satir in enumerate(satirlar, 2):
        for col_idx, deger in enumerate(satir, 1):
            hucre = ws.cell(row=row_idx, column=col_idx, value=deger)
            if row_idx % 2 == 0:
                hucre.fill = zebra_doldur

    # Sütun genişliği
    for col_idx, sutun in enumerate(sutunlar, 1):
        max_genislik = len(str(sutun))
        for satir in satirlar:
            if col_idx - 1 < len(satir):
                max_genislik = max(max_genislik, len(str(satir[col_idx - 1] or "")))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_genislik + 4, 50)

    wb.save(kayit_yolu)
