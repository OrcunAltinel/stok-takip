"""Tosun Altınel Fiyat Teklifi Excel dosyasından ilk ürün girişi.

Çalıştırma:
    py araclar/excel_ice_aktar.py

Excel'deki her geçerli satır için:
  - ürün yoksa `urun_servisi.urun_ekle` ile oluşturulur, varsa fiyatları/rapa
    kodu `urun_servisi.urun_guncelle` ile güncellenir,
  - `stok_servisi.stok_girisi_kaydet` ile GIRIS tipinde ilk stok girişi
    yapılır (mükerrer çalıştırmada bu adım atlanır).

Not: `veritabani/baglanti.py::get_session()` her servis çağrısında kendi
transaction'ını commit eder (bkz. CLAUDE.md "Session Yönetimi"). Bu yüzden
1.084 satırlık import'u tek bir DB transaction'ında toplamak, servis
katmanını (istenen mimari) bozmadan mümkün değil. Bunun yerine:
  1) import öncesi tam DB yedeği alınır,
  2) her satır ayrı ayrı işlenir ve hata alan satır atlanıp raporlanır
     (diğer satırlar etkilenmez),
  3) tüm adımlar idempotenttir — betik tekrar çalıştırılınca zaten
     aktarılmış ürün/stok girişleri tekrar oluşturulmaz.
Bu, satır bazlı bir hatanın diğer 1000+ satırı geri almasındansa daha
kullanışlı ve pratikte aynı güvenliği (yedekten dönebilme + tekrar
çalıştırılabilirlik) sağlıyor.
"""

import random
import sys
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from sqlalchemy import func

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.stdout.reconfigure(encoding="utf-8")

from servisler import stok_servisi, urun_servisi, yedek_servisi
from veritabani.baglanti import engine, get_session, veritabani_olustur
from veritabani.migrasyonlar import tum_migrasyonlari_calistir
from veritabani.modeller import Admin, StokHareketi, Tedarikci, Urun

EXCEL_YOLU = Path(__file__).resolve().parent.parent / "Tosun Altınel  Fiyat Teklifi.xlsx"
SAYFA_ADI = "Sayfa1"
TEDARIKCI_ADI = "Rapa Limited"
GIRIS_ACIKLAMASI = "İlk ürün girişi — Tosun Altınel Fiyat Teklifi"

# Excel sütun sırası sabit — başlık metnindeki Türkçe karakter/boşluk
# farklılıklarına karşı konumsal (0 tabanlı) indeks kullanılıyor.
IDX_RAPA_KODU = 0
IDX_STOK_KODU = 1
IDX_STOK_ADI = 2
IDX_GIDECEK_STOK = 3
IDX_KDVLI_FIYAT = 4
IDX_INDIRIMLI_FIYAT = 5


def _ondalik(deger) -> Decimal:
    d = Decimal(str(deger))
    return d.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def admin_id_bul() -> int:
    with get_session() as session:
        admin = session.query(Admin).filter_by(kullanici_adi="admin", aktif=True).first()
        if not admin:
            admin = session.query(Admin).filter_by(aktif=True).order_by(Admin.id).first()
        if not admin:
            raise RuntimeError("Aktif admin bulunamadı; içe aktarma için en az bir admin hesabı gerekli.")
        admin_id = admin.id
    return admin_id


def tedarikci_id_bul_veya_olustur() -> int:
    with get_session() as session:
        t = session.query(Tedarikci).filter_by(firma_adi=TEDARIKCI_ADI).first()
        if not t:
            t = Tedarikci(firma_adi=TEDARIKCI_ADI, aktif=True)
            session.add(t)
            session.flush()
        tedarikci_id = t.id
    return tedarikci_id


def excel_satirlarini_oku() -> list[tuple[int, tuple]]:
    if not EXCEL_YOLU.exists():
        raise FileNotFoundError(f"Excel dosyası bulunamadı: {EXCEL_YOLU}")
    wb = openpyxl.load_workbook(EXCEL_YOLU, data_only=True, read_only=True)
    if SAYFA_ADI not in wb.sheetnames:
        raise ValueError(f"'{SAYFA_ADI}' sayfası bulunamadı. Mevcut sayfalar: {wb.sheetnames}")
    ws = wb[SAYFA_ADI]
    satirlar = list(enumerate(ws.iter_rows(min_row=2, values_only=True), start=2))
    wb.close()
    return satirlar


def satiri_isle(satir_no: int, row: tuple) -> dict:
    stok_kodu_ham = row[IDX_STOK_KODU] if len(row) > IDX_STOK_KODU else None
    if stok_kodu_ham is None or str(stok_kodu_ham).strip() == "":
        return {"durum": "bos"}

    stok_kodu = str(stok_kodu_ham).strip().upper()
    stok_adi = str(row[IDX_STOK_ADI] or "").strip()

    rapa_ham = row[IDX_RAPA_KODU]
    rapa_kodu = str(rapa_ham).strip() if rapa_ham not in (None, "") else None

    miktar_ham = row[IDX_GIDECEK_STOK]
    try:
        miktar = int(str(miktar_ham).strip())
    except (ValueError, TypeError):
        return {"durum": "hata", "sebep": f"geçersiz miktar: {miktar_ham!r}"}

    try:
        kdvli_fiyat = _ondalik(row[IDX_KDVLI_FIYAT])
        satis_fiyati = _ondalik(row[IDX_INDIRIMLI_FIYAT])
    except (InvalidOperation, TypeError):
        return {"durum": "hata", "sebep": "fiyat çevrilemedi"}

    return {
        "durum": "gecerli",
        "satir_no": satir_no,
        "stok_kodu": stok_kodu,
        "stok_adi": stok_adi,
        "rapa_kodu": rapa_kodu,
        "miktar": miktar,
        "kdvli_fiyat": kdvli_fiyat,
        "satis_fiyati": satis_fiyati,
    }


def daha_once_aktarilmis_mi(urun_id: int) -> bool:
    with get_session() as session:
        var = (
            session.query(StokHareketi)
            .filter_by(urun_id=urun_id, hareket_tipi="GIRIS", aciklama=GIRIS_ACIKLAMASI)
            .first()
        )
        return var is not None


def ice_aktar() -> None:
    print("=== TOSUN ALTINEL FİYAT TEKLİFİ İÇE AKTARMA ===\n")

    veritabani_olustur()
    tum_migrasyonlari_calistir(engine)

    print("Veritabanı yedekleniyor...")
    yedek_yolu = yedek_servisi.yedek_al(etiket="import_oncesi")
    print(f"Yedek alındı: {yedek_yolu}\n")

    admin_id = admin_id_bul()
    tedarikci_id = tedarikci_id_bul_veya_olustur()
    print(f"Tedarikçi: {TEDARIKCI_ADI} (id={tedarikci_id})")
    print(f"Admin: id={admin_id}\n")

    satirlar = excel_satirlarini_oku()
    okunan = len(satirlar)

    bos_satir = 0
    hatali_satirlar: list[tuple[int, str]] = []
    yeni_urun = 0
    guncellenen_urun = 0
    olusturulan_giris = 0
    onceden_aktarilmis = 0
    toplam_giris_miktari = 0
    beklenen_toplam_miktar = 0
    gecerli_satirlar: list[dict] = []

    for satir_no, row in satirlar:
        sonuc = satiri_isle(satir_no, row)

        if sonuc["durum"] == "bos":
            bos_satir += 1
            continue

        if sonuc["durum"] == "hata":
            hatali_satirlar.append((satir_no, sonuc["sebep"]))
            continue

        gecerli_satirlar.append(sonuc)
        beklenen_toplam_miktar += sonuc["miktar"]

        try:
            urun = urun_servisi.urun_bul_kod(sonuc["stok_kodu"])
            if urun is None:
                urun = urun_servisi.urun_ekle(
                    urun_kodu=sonuc["stok_kodu"],
                    urun_adi=sonuc["stok_adi"] or sonuc["stok_kodu"],
                    rapa_kodu=sonuc["rapa_kodu"],
                    birim="adet",
                    alis_fiyati=sonuc["satis_fiyati"],
                    satis_fiyati=sonuc["satis_fiyati"],
                    kdvli_fiyat=sonuc["kdvli_fiyat"],
                    kritik_stok_seviyesi=Decimal("5"),
                )
                yeni_urun += 1
            else:
                urun_servisi.urun_guncelle(
                    urun.id,
                    rapa_kodu=sonuc["rapa_kodu"],
                    satis_fiyati=sonuc["satis_fiyati"],
                    kdvli_fiyat=sonuc["kdvli_fiyat"],
                )
                guncellenen_urun += 1

            if daha_once_aktarilmis_mi(urun.id):
                onceden_aktarilmis += 1
                continue

            stok_servisi.stok_girisi_kaydet(
                urun_id=urun.id,
                miktar=Decimal(sonuc["miktar"]),
                birim_fiyat=sonuc["satis_fiyati"],
                admin_id=admin_id,
                tedarikci_id=tedarikci_id,
                aciklama=GIRIS_ACIKLAMASI,
            )
            olusturulan_giris += 1
            toplam_giris_miktari += sonuc["miktar"]
        except Exception as e:
            hatali_satirlar.append((satir_no, str(e)))

    print("=== İÇE AKTARMA RAPORU ===")
    print(f"Okunan satır          : {okunan}")
    print(f"Boş/atlanan satır     : {bos_satir}")
    print(f"Eklenen yeni ürün     : {yeni_urun}")
    print(f"Güncellenen ürün      : {guncellenen_urun}")
    print(f"Oluşturulan GIRIS     : {olusturulan_giris}")
    print(f"Önceden aktarılmış    : {onceden_aktarilmis}")
    print(f"Bu çalıştırmada giren adet : {toplam_giris_miktari}")
    print(f"Excel'den hesaplanan beklenen toplam adet : {beklenen_toplam_miktar}")
    if hatali_satirlar:
        print(f"Hatalı satırlar       : {hatali_satirlar}")
    else:
        print("Hatalı satırlar       : yok")

    dogrula(gecerli_satirlar, beklenen_toplam_miktar)


def dogrula(gecerli_satirlar: list[dict], beklenen_toplam_miktar: int) -> None:
    print("\n=== DOĞRULAMA ===")

    with get_session() as session:
        aktif_urun_sayisi = session.query(Urun).filter_by(aktif=True).count()
        toplam_miktar = (
            session.query(func.sum(StokHareketi.miktar))
            .filter(StokHareketi.hareket_tipi == "GIRIS", StokHareketi.aciklama == GIRIS_ACIKLAMASI)
            .scalar()
        ) or Decimal("0")

    beklenen_urun_sayisi = len(gecerli_satirlar)
    print(f"1) Aktif ürün sayısı: {aktif_urun_sayisi} (beklenen >= {beklenen_urun_sayisi})")
    if aktif_urun_sayisi < beklenen_urun_sayisi:
        print("   UYARI: aktif ürün sayısı beklenenden az!")

    print(f"2) Bu import'a ait toplam GIRIS miktarı: {toplam_miktar} (Excel'den beklenen: {beklenen_toplam_miktar})")
    if Decimal(toplam_miktar) != Decimal(beklenen_toplam_miktar):
        print("   UYARI: toplam miktar Excel'den hesaplanan beklenen değerle uyuşmuyor!")

    print("3) Rastgele 5 ürün karşılaştırması (Excel vs. veritabanı):")
    ornekler = random.sample(gecerli_satirlar, min(5, len(gecerli_satirlar)))
    for veri in ornekler:
        urun = urun_servisi.urun_bul_kod(veri["stok_kodu"])
        if urun is None:
            print(f"   {veri['stok_kodu']}: veritabanında bulunamadı!")
            continue
        print(
            f"   {veri['stok_kodu']}: "
            f"satış Excel={veri['satis_fiyati']} / DB={urun.satis_fiyati}  |  "
            f"KDV'li Excel={veri['kdvli_fiyat']} / DB={urun.kdvli_fiyat}"
        )

    print("4) stok_miktarı tutarlılık örneklemi (10 ürün):")
    ornekler_2 = random.sample(gecerli_satirlar, min(10, len(gecerli_satirlar)))
    with get_session() as session:
        for veri in ornekler_2:
            urun = session.query(Urun).filter_by(urun_kodu=veri["stok_kodu"]).first()
            if urun is None:
                continue
            hareketler = session.query(StokHareketi).filter_by(urun_id=urun.id).all()
            if any(h.hareket_tipi == "SAYIM_DUZELTME" for h in hareketler):
                print(f"   {veri['stok_kodu']}: atlandı (sayım düzeltmesi mevcut, basit toplamla doğrulanamaz)")
                continue
            hesaplanan = sum(
                (h.miktar if h.hareket_tipi in ("GIRIS", "IADE_GIRIS") else -h.miktar)
                for h in hareketler
            )
            durum = "OK" if Decimal(str(hesaplanan)) == Decimal(str(urun.stok_miktari)) else "UYUŞMUYOR"
            print(f"   {veri['stok_kodu']}: hesaplanan={hesaplanan} / stok_miktari={urun.stok_miktari}  [{durum}]")


if __name__ == "__main__":
    ice_aktar()
